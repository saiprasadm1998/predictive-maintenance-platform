"""Hierarchy and tag import from CSV or Excel (UC-03, FR-REG-03).

Validate first, then commit atomically: the whole file is parsed and checked, every
problem is reported with its sheet, line and column, and nothing is written unless the
file is completely valid. Re-importing the same file changes nothing (idempotent).

Sheets (Excel) or files (CSV, named after the sheet):
    asset_types  code, name, description
    roles        asset_type, code, name, unit
    assets       code, name, level, parent, criticality, asset_type, description
    tags         source, source_name, asset, description, unit, range_lo, range_hi,
                 interval_s, role
"""

from __future__ import annotations

import csv
import io
from collections import Counter
from collections.abc import Iterable, Mapping
from dataclasses import dataclass, field
from functools import partial
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from pmp.core import Actor
from pmp.registry.domain import (
    AssetLevel,
    Criticality,
    RegistryError,
    TagSpec,
    check_code,
    check_criticality,
    check_name,
    check_parent_level,
    check_role_code,
)
from pmp.registry.models import Asset, AssetType, TypeRole
from pmp.registry.service import (
    AssetSpec,
    Change,
    RoleSpec,
    upsert_asset,
    upsert_asset_type,
    upsert_role,
    upsert_tag,
)

SHEETS: dict[str, tuple[set[str], set[str]]] = {
    # sheet: (required columns, optional columns)
    "asset_types": ({"code", "name"}, {"description"}),
    "roles": ({"asset_type", "code", "name"}, {"unit"}),
    "assets": ({"code", "name", "level"}, {"parent", "criticality", "asset_type", "description"}),
    "tags": (
        {"source", "source_name", "asset", "unit", "range_lo", "range_hi", "interval_s"},
        {"description", "role"},
    ),
}
MAX_ROWS = 100_000


@dataclass(frozen=True)
class Row:
    sheet: str
    line: int  # line number as the user sees it (header is line 1)
    values: Mapping[str, str]

    def get(self, column: str) -> str:
        return self.values.get(column, "").strip()


@dataclass(frozen=True)
class Problem:
    sheet: str
    line: int
    column: str
    message: str

    def __str__(self) -> str:
        where = f"{self.sheet}, line {self.line}" if self.line else self.sheet
        if self.column:
            where += f", column {self.column}"
        return f"{where}: {self.message}"


@dataclass
class ImportReport:
    problems: list[Problem] = field(default_factory=list)
    counts: dict[str, Counter[str]] = field(default_factory=dict)
    committed: bool = False

    @property
    def ok(self) -> bool:
        return not self.problems

    def count(self, sheet: str, change: Change) -> None:
        self.counts.setdefault(sheet, Counter())[change.value] += 1

    def summary(self) -> str:
        if self.problems:
            lines = [f"Import rejected: {len(self.problems)} problem(s); nothing was changed."]
            lines += [f"  {p}" for p in self.problems]
            return "\n".join(lines)
        lines = ["Import committed." if self.committed else "Import is valid (dry run)."]
        for sheet in SHEETS:
            if sheet in self.counts:
                c = self.counts[sheet]
                lines.append(
                    f"  {sheet}: {c['created']} created, {c['updated']} updated, "
                    f"{c['unchanged']} unchanged"
                )
        return "\n".join(lines)


# ------------------------------------------------------------------ reading files


def _normalise_header(cell: object) -> str:
    return str(cell or "").strip().lower().replace(" ", "_")


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _rows_from_table(
    sheet: str, header: list[object], body: Iterable[Iterable[object]]
) -> list[Row]:
    columns = [_normalise_header(h) for h in header]
    rows: list[Row] = []
    for offset, raw in enumerate(body, start=2):
        cells = [_cell_text(v) for v in raw]
        if not any(c.strip() for c in cells):
            continue  # skip blank lines
        values = {col: cells[i] if i < len(cells) else "" for i, col in enumerate(columns) if col}
        rows.append(Row(sheet, offset, values))
    return rows


def read_csv(sheet: str, data: bytes) -> list[Row]:
    text = data.decode("utf-8-sig")  # Excel adds a byte-order mark
    try:
        dialect: Any = csv.Sniffer().sniff(text.splitlines()[0] if text else ",", ",;\t")
    except csv.Error:
        dialect = csv.excel
    table = list(csv.reader(io.StringIO(text), dialect))
    if not table:
        return []
    return _rows_from_table(sheet, list(table[0]), table[1:])


def read_workbook(data: bytes) -> dict[str, list[Row]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets: dict[str, list[Row]] = {}
    for worksheet in workbook.worksheets:
        name = _normalise_header(worksheet.title)
        if name not in SHEETS:
            continue
        table = list(worksheet.iter_rows(values_only=True))
        if table:
            sheets[name] = _rows_from_table(name, list(table[0]), table[1:])
    workbook.close()
    return sheets


def read_files(paths: Iterable[Path]) -> tuple[dict[str, list[Row]], list[Problem]]:
    """Read .xlsx workbooks and/or .csv files named after their sheet (e.g. assets.csv)."""
    sheets: dict[str, list[Row]] = {}
    problems: list[Problem] = []
    for path in paths:
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            found = read_workbook(path.read_bytes())
        elif suffix == ".csv":
            name = _normalise_header(path.stem)
            if name not in SHEETS:
                problems.append(
                    Problem(
                        path.name,
                        0,
                        "",
                        f"unknown sheet; name the file one of: "
                        f"{', '.join(f'{s}.csv' for s in SHEETS)}",
                    )
                )
                continue
            found = {name: read_csv(name, path.read_bytes())}
        else:
            problems.append(Problem(path.name, 0, "", "only .csv and .xlsx files are supported"))
            continue
        for name, rows in found.items():
            if name in sheets:
                problems.append(Problem(name, 0, "", f"sheet given more than once ({path.name})"))
            sheets[name] = rows
    return sheets, problems


# ------------------------------------------------------------------ validation


@dataclass(frozen=True)
class _Plan:
    asset_types: list[tuple[Row, str, str, str]]
    roles: list[tuple[Row, str, RoleSpec]]
    assets: list[tuple[Row, AssetSpec]]
    tags: list[tuple[Row, TagSpec, str | None]]


def _number(row: Row, column: str, problems: list[Problem]) -> float | None:
    raw = row.get(column)
    if not raw:
        problems.append(Problem(row.sheet, row.line, column, "a number is required"))
        return None
    if "," in raw and "." not in raw:
        raw = raw.replace(",", ".")  # decimal comma
    try:
        return float(raw)
    except ValueError:
        problems.append(
            Problem(row.sheet, row.line, column, f"{row.get(column)!r} is not a number")
        )
        return None


def _integer(row: Row, column: str, problems: list[Problem]) -> int | None:
    value = _number(row, column, problems)
    if value is None:
        return None
    if not value.is_integer():
        problems.append(Problem(row.sheet, row.line, column, "a whole number is required"))
        return None
    return int(value)


def _try(row: Row, column: str, problems: list[Problem], check: Any, *args: Any) -> Any:
    try:
        return check(*args)
    except RegistryError as error:
        problems.append(Problem(row.sheet, row.line, column, str(error)))
        return None


def _check_columns(sheets: Mapping[str, list[Row]], problems: list[Problem]) -> None:
    for sheet, rows in sheets.items():
        required, optional = SHEETS[sheet]
        if not rows:
            continue
        present = set(rows[0].values)
        missing = sorted(required - present)
        unknown = sorted(present - required - optional)
        if missing:
            problems.append(Problem(sheet, 1, "", f"missing column(s): {', '.join(missing)}"))
        if unknown:
            problems.append(Problem(sheet, 1, "", f"unknown column(s): {', '.join(unknown)}"))
        if len(rows) > MAX_ROWS:
            problems.append(Problem(sheet, 0, "", f"more than {MAX_ROWS} rows"))


def _duplicates(rows: list[Row], key: Any, what: str, problems: list[Problem]) -> None:
    first_line: dict[Any, int] = {}
    for row in rows:
        k = key(row)
        if not k or (isinstance(k, tuple) and not all(k)):
            continue
        if k in first_line:
            problems.append(
                Problem(
                    row.sheet,
                    row.line,
                    "",
                    f"duplicate {what} (first given on line {first_line[k]})",
                )
            )
        else:
            first_line[k] = row.line


def validate(session: Session, sheets: Mapping[str, list[Row]]) -> tuple[_Plan, list[Problem]]:
    """Check every row against the file itself and the database. Writes nothing."""
    problems: list[Problem] = []
    _check_columns(sheets, problems)
    if problems:
        return _Plan([], [], [], []), problems

    type_rows = sheets.get("asset_types", [])
    role_rows = sheets.get("roles", [])
    asset_rows = sheets.get("assets", [])
    tag_rows = sheets.get("tags", [])

    _duplicates(type_rows, lambda r: r.get("code"), "asset type code", problems)
    _duplicates(role_rows, lambda r: (r.get("asset_type"), r.get("code")), "role", problems)
    _duplicates(asset_rows, lambda r: r.get("code"), "asset code", problems)
    _duplicates(
        tag_rows,
        lambda r: (r.get("source"), r.get("source_name")),
        "tag (source, source_name)",
        problems,
    )

    # Asset types
    db_types = {code for (code,) in session.execute(select(AssetType.code))}
    file_types: set[str] = set()
    plan_types = []
    for row in type_rows:
        code = _try(row, "code", problems, check_code, row.get("code"), "Asset type code")
        name = _try(row, "name", problems, check_name, row.get("name"), "Asset type name")
        if code and name:
            file_types.add(code)
            plan_types.append((row, code, name, row.get("description")))
    known_types = db_types | file_types

    # Roles
    db_roles = {
        (t, r)
        for t, r in session.execute(
            select(AssetType.code, TypeRole.code).join(
                TypeRole, TypeRole.asset_type_id == AssetType.id
            )
        )
    }
    plan_roles = []
    for row in role_rows:
        type_code = row.get("asset_type")
        if type_code not in known_types:
            problems.append(
                Problem(
                    row.sheet, row.line, "asset_type", f"asset type {type_code!r} is not defined"
                )
            )
            continue
        code = _try(row, "code", problems, check_role_code, row.get("code"))
        name = _try(row, "name", problems, check_name, row.get("name"), "Role name")
        if code and name:
            plan_roles.append((row, type_code, RoleSpec(code, name, row.get("unit"))))
    known_roles = db_roles | {(t, spec.code) for _, t, spec in plan_roles}

    # Assets: parents may be in the file (any order) or already in the database
    db_levels = {
        code: AssetLevel(level) for code, level in session.execute(select(Asset.code, Asset.level))
    }
    file_assets: dict[str, tuple[Row, AssetLevel, str | None]] = {}
    specs: dict[str, tuple[Row, AssetSpec]] = {}
    for row in asset_rows:
        code = _try(row, "code", problems, check_code, row.get("code"), "Asset code")
        name = _try(row, "name", problems, check_name, row.get("name"), "Asset name")
        level = _try(row, "level", problems, AssetLevel.parse, row.get("level"))
        criticality = (
            _try(row, "criticality", problems, Criticality.parse, row.get("criticality"))
            if row.get("criticality")
            else None
        )
        asset_type = row.get("asset_type") or None
        if asset_type and asset_type not in known_types:
            problems.append(
                Problem(
                    row.sheet, row.line, "asset_type", f"asset type {asset_type!r} is not defined"
                )
            )
        if level is not None:
            _try(row, "criticality", problems, check_criticality, level, criticality)
            if code and code in db_levels and db_levels[code] != level:
                problems.append(
                    Problem(
                        row.sheet,
                        row.line,
                        "level",
                        f"asset {code!r} already exists as a "
                        f"{db_levels[code].label}; its level cannot change",
                    )
                )
        if code and name and level is not None:
            parent = row.get("parent") or None
            file_assets[code] = (row, level, parent)
            specs[code] = (
                row,
                AssetSpec(
                    code, name, level, parent, criticality, asset_type, row.get("description")
                ),
            )

    def level_of(code: str) -> AssetLevel | None:
        if code in file_assets:
            return file_assets[code][1]
        return db_levels.get(code)

    for row, level, parent in file_assets.values():
        if parent is None:
            _try(row, "level", problems, check_parent_level, level, None)
            continue
        parent_level = level_of(parent)
        if parent_level is None:
            problems.append(
                Problem(
                    row.sheet,
                    row.line,
                    "parent",
                    f"parent {parent!r} is not in the file or the registry",
                )
            )
            continue
        _try(row, "parent", problems, check_parent_level, level, parent_level)

    ordered_assets = _parents_first(specs, file_assets, problems)

    # Tags
    known_assets = set(db_levels) | set(file_assets)
    asset_type_of: dict[str, str | None] = {
        code: spec.asset_type_code for code, (_, spec) in specs.items()
    }
    for code, type_code in session.execute(
        select(Asset.code, AssetType.code).outerjoin(AssetType, AssetType.id == Asset.asset_type_id)
    ):
        asset_type_of.setdefault(code, type_code)
    plan_tags = []
    role_holders: dict[tuple[str, str], int] = {}
    for row in tag_rows:
        asset_code = row.get("asset")
        if asset_code not in known_assets:
            problems.append(
                Problem(
                    row.sheet,
                    row.line,
                    "asset",
                    f"asset {asset_code!r} is not in the file or the registry",
                )
            )
        range_lo = _number(row, "range_lo", problems)
        range_hi = _number(row, "range_hi", problems)
        interval = _integer(row, "interval_s", problems)
        role = row.get("role") or None
        if role:
            role_type = asset_type_of.get(asset_code)
            if role_type is None:
                problems.append(
                    Problem(
                        row.sheet,
                        row.line,
                        "role",
                        f"asset {asset_code!r} has no asset type, so the tag cannot have a role",
                    )
                )
            elif (role_type, role) not in known_roles:
                problems.append(
                    Problem(
                        row.sheet,
                        row.line,
                        "role",
                        f"asset type {role_type!r} has no role {role!r}",
                    )
                )
            elif (asset_code, role) in role_holders:
                problems.append(
                    Problem(
                        row.sheet,
                        row.line,
                        "role",
                        f"role {role!r} on {asset_code!r} is already given to "
                        f"the tag on line {role_holders[(asset_code, role)]}",
                    )
                )
            else:
                role_holders[(asset_code, role)] = row.line
        if range_lo is None or range_hi is None or interval is None:
            continue
        spec = _try(
            row,
            "",
            problems,
            TagSpec(
                source=row.get("source"),
                source_name=row.get("source_name"),
                asset_code=asset_code,
                unit=row.get("unit"),
                range_lo=range_lo,
                range_hi=range_hi,
                interval_s=interval,
                description=row.get("description"),
            ).validated,
        )
        if spec is not None:
            plan_tags.append((row, spec, role))

    problems.sort(key=lambda p: (list(SHEETS).index(p.sheet) if p.sheet in SHEETS else 99, p.line))
    return _Plan(plan_types, plan_roles, ordered_assets, plan_tags), problems


def _parents_first(
    specs: Mapping[str, tuple[Row, AssetSpec]],
    file_assets: Mapping[str, tuple[Row, AssetLevel, str | None]],
    problems: list[Problem],
) -> list[tuple[Row, AssetSpec]]:
    """Order assets so every parent in the file comes before its children; detect cycles."""
    ordered: list[tuple[Row, AssetSpec]] = []
    state: dict[str, str] = {}

    def visit(code: str, trail: list[str]) -> None:
        if state.get(code) == "done":
            return
        if state.get(code) == "visiting":
            row = file_assets[code][0]
            problems.append(
                Problem(
                    row.sheet,
                    row.line,
                    "parent",
                    f"circular parent chain: {' > '.join([*trail, code])}",
                )
            )
            return
        state[code] = "visiting"
        parent = file_assets[code][2]
        if parent is not None and parent in file_assets:
            visit(parent, [*trail, code])
        state[code] = "done"
        if code in specs:
            ordered.append(specs[code])

    for code in file_assets:
        visit(code, [])
    return ordered


# ------------------------------------------------------------------ import


def import_registry(
    session: Session, actor: Actor, sheets: Mapping[str, list[Row]], *, dry_run: bool = False
) -> ImportReport:
    """Validate, then apply everything in the caller's transaction.

    On any problem nothing is written. The caller's transaction makes the import atomic:
    if anything fails while applying, the caller rolls back.
    """
    report = ImportReport()
    plan, problems = validate(session, sheets)
    if problems:
        report.problems = problems
        return report
    if dry_run:
        return report

    for row, code, name, description in plan.asset_types:
        _apply(
            report,
            row,
            partial(
                upsert_asset_type, session, actor, code=code, name=name, description=description
            ),
        )
    for row, type_code, role in plan.roles:
        _apply(report, row, partial(upsert_role, session, actor, type_code, role))
    for row, spec in plan.assets:
        _apply(report, row, partial(upsert_asset, session, actor, spec))
    for row, tag_spec, role_code in plan.tags:
        _apply(report, row, partial(upsert_tag, session, actor, tag_spec, role_code=role_code))

    if report.problems:
        raise ImportFailedError(report)
    report.committed = True
    return report


class ImportFailedError(Exception):
    """Raised when a row fails while applying, so the caller's transaction rolls back."""

    def __init__(self, report: ImportReport) -> None:
        super().__init__(report.summary())
        self.report = report


def _apply(report: ImportReport, row: Row, action: Any) -> None:
    if report.problems:
        return  # stop after the first failure; the transaction will be rolled back
    try:
        change: Change = action()
    except RegistryError as error:
        report.problems.append(Problem(row.sheet, row.line, "", str(error)))
        return
    report.count(row.sheet, change)


def import_files(
    session: Session, actor: Actor, paths: Iterable[Path], *, dry_run: bool = False
) -> ImportReport:
    sheets, problems = read_files(paths)
    if problems:
        return ImportReport(problems=problems)
    return import_registry(session, actor, sheets, dry_run=dry_run)


__all__ = [
    "ImportFailedError",
    "ImportReport",
    "Problem",
    "Row",
    "import_files",
    "import_registry",
    "read_csv",
    "read_files",
    "read_workbook",
    "validate",
]
